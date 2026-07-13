
# ########################################Folder wise Extraction################
# """
# ALL-IN-ONE extractor for the HRSG control narrative (.docx):

#   Output 1: the block diagram of section 4.1.1 as PNG (via PowerPoint)
#   Output 2: the not-struck-out diagram text  -> .txt
#   Output 3: tables of all sections referenced by the diagram -> .xlsx
#             with PIN MAPPING: every table row label is mapped to its
#             MarkVIe actuator block pin name (extra "Pin" column)
#             plus a dedicated "Pins" sheet listing all mapped pins

#   Folder structure of the output:
#     OUTPUT_DIR\\overall\\      -> image + text + combined Excel
#     OUTPUT_DIR\\<Mnemonic>\\   -> <Mnemonic>.xlsx (its sheet + Pins sheet)

# Requires:  pip install python-docx openpyxl pillow pywin32
# Usage   :  edit CONFIG, then run  python Image_extract.py
# """

# print("script started")

# import io
# import re
# import traceback
# import xml.etree.ElementTree as ET
# import zipfile
# from pathlib import Path

# import docx
# from docx.table import Table
# from docx.text.paragraph import Paragraph

# # ============================================================
# #                          CONFIG
# # ============================================================
# DOC_PATH = r"C:\Users\2023988\Desktop\HRSG\Control Narrative and XML\YTL-50-E-LBA50---EN-LO-001 - 1GP052419 - A-en-1GP052419AT03-en.docx"

# TARGET_SECTION = "4.1.1"          # section containing the block diagram

# OUTPUT_DIR = r"C:\Users\2023988\Desktop\HRSG\code\output\out_image_1"
# IMAGE_NAME = "section_4_1_1_diagram"          # -> .emf + .png
# TEXT_NAME  = "section_4_1_1_diagram_text.txt"
# XLSX_NAME  = "referenced_section_tables.xlsx"

# PNG_SCALE = 4                     # enlarge before export = sharper PNG

# # Narrative table row label  ->  MarkVIe block pin name
# # (add more entries here any time - matching is case-insensitive)
# PIN_MAP = {
#     # group (ON/OFF) pins
#     "set auto mode":      "AU_SEL",
#     "permit on":          "ON_PMT",
#     "auto command on":    "AU_ON",
#     "force on":           "ON_FRC",
#     "permit off":         "OFF_PMT",
#     "auto command off":   "AU_OFF",
#     "force off":          "OFF_FRC",
#     # valve (OPEN/CLOSE) pins
#     "permit open":        "OP_PMT",
#     "auto command open":  "AU_OP",
#     "force open":         "OP_FRC",
#     "override open":      "OP_OVR",
#     "permit close":       "CL_PMT",
#     "auto command close": "AU_CL",
#     "force close":        "CL_FRC",
#     "override close":     "CL_OVR",
#     "auto permit":        "AU_PMT",
#     "force condition":    "CMD_FRC",
#     "force command":      "CMD_FRC",
#     # control valve pins
#     "modulate permit":    "MOD_PMT",
#     "modulate value":     "RSP",
#     "force value":        "FV",
#     "auto setpoint":      "AU_SETP",
# }
# # ============================================================

# VNS = "{http://schemas.microsoft.com/office/visio/2012/main}"
# WNS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# HEADING_RX = re.compile(r"^Heading (\d)$")
# SECREF_RX = re.compile(r"\[\s*\u00a7\s*([\d.]+)\s*\]")


# def full_text(el):
#     """Text including tracked-change deletions (original document)."""
#     return "".join(n.text or "" for n in el.iter()
#                    if n.tag in (WNS + "t", WNS + "delText"))


# def live_text(el):
#     """Text as it reads after accepting tracked changes."""
#     return "".join(n.text or "" for n in el.iter() if n.tag == WNS + "t")


# def pin_for(label):
#     """Map a narrative row label to a block pin name ('' if unknown)."""
#     return PIN_MAP.get(" ".join(label.lower().split()), "")


# # ---------- locate diagram objects in the target section ----------

# def find_section_objects(doc):
#     """Return (image_rid, ole_rid) of the diagram in TARGET_SECTION."""
#     counters = [0] * 9
#     depth = TARGET_SECTION.count(".") + 1
#     inside = False
#     image_rid = ole_rid = None

#     for par in doc.paragraphs:
#         m = HEADING_RX.match(par.style.name)
#         if m and full_text(par._p).strip():
#             lvl = int(m.group(1))
#             counters[lvl - 1] += 1
#             for i in range(lvl, 9):
#                 counters[i] = 0
#             number = ".".join(str(c) for c in counters[:lvl])
#             if inside and lvl <= depth:
#                 break
#             if number == TARGET_SECTION:
#                 inside = True
#                 print("Found section %s: %s" % (number, par.text.strip()))
#             continue
#         if inside:
#             xml = par._p.xml
#             if image_rid is None:
#                 m2 = re.search(r'<v:imagedata[^>]*r:id="(rId\d+)"', xml)
#                 if m2:
#                     image_rid = m2.group(1)
#             if ole_rid is None:
#                 m3 = re.search(r'<o:OLEObject[^>]*r:id="(rId\d+)"', xml)
#                 if m3:
#                     ole_rid = m3.group(1)
#             if image_rid and ole_rid:
#                 break

#     if not (image_rid or ole_rid):
#         raise SystemExit("No diagram found in section %s." % TARGET_SECTION)
#     return image_rid, ole_rid


# # ---------- output 1: image ----------

# def save_image(doc, image_rid, out_dir):
#     part = doc.part.related_parts[image_rid]
#     ext = Path(str(part.partname)).suffix
#     raw = out_dir / (IMAGE_NAME + ext)
#     raw.write_bytes(part.blob)
#     print("Saved original image: %s" % raw)

#     if ext.lower() not in (".emf", ".wmf"):
#         return

#     png = out_dir / (IMAGE_NAME + ".png")
#     try:
#         import win32com.client
#         print("Starting PowerPoint for EMF -> PNG conversion ...")
#         ppt = win32com.client.Dispatch("PowerPoint.Application")
#         pres = ppt.Presentations.Add(WithWindow=0)
#         try:
#             slide = pres.Slides.Add(1, 12)
#             shape = slide.Shapes.AddPicture(
#                 FileName=str(raw), LinkToFile=0, SaveWithDocument=-1,
#                 Left=0, Top=0)
#             shape.ScaleWidth(PNG_SCALE, 1)
#             shape.ScaleHeight(PNG_SCALE, 1)
#             shape.Export(str(png), 2)          # 2 = PNG
#         finally:
#             pres.Close()
#             ppt.Quit()

#         try:
#             from PIL import Image, ImageChops
#             img = Image.open(png).convert("RGB")
#             bg = Image.new("RGB", img.size, (255, 255, 255))
#             bbox = ImageChops.difference(img, bg).getbbox()
#             if bbox:
#                 img.crop(bbox).save(png)
#         except ImportError:
#             pass
#         print("Saved PNG           : %s" % png)
#     except Exception as e:
#         print("[!] PNG conversion failed (%s); original %s kept." % (e, ext))


# # ---------- output 2: diagram text (not struck out) ----------

# def diagram_kept_texts(vsdx_bytes):
#     zf = zipfile.ZipFile(io.BytesIO(vsdx_bytes))
#     page = [n for n in zf.namelist()
#             if re.match(r"visio/pages/page\d+\.xml$", n)][0]
#     root = ET.fromstring(zf.read(page))

#     out = []
#     for shape in root.iter(VNS + "Shape"):
#         text_el = shape.find(VNS + "Text")
#         if text_el is None:
#             continue
#         struck_rows = {}
#         for sec in shape.findall(VNS + "Section"):
#             if sec.get("N") == "Character":
#                 for row in sec.findall(VNS + "Row"):
#                     ix = int(row.get("IX", "0"))
#                     s = any(c.get("N") == "Strikethru" and c.get("V") == "1"
#                             for c in row.findall(VNS + "Cell"))
#                     struck_rows[ix] = s
#         kept = []
#         state = {"cur": 0}

#         def add(t):
#             if t and not struck_rows.get(state["cur"], False):
#                 kept.append(t)

#         add(text_el.text)
#         for ch in text_el:
#             if ch.tag == VNS + "cp":
#                 state["cur"] = int(ch.get("IX", "0"))
#             add(ch.tail)
#         txt = " ".join("".join(kept).split())
#         if txt:
#             out.append(txt)
#     return out


# def pair_refs_with_names(texts):
#     refs = {}
#     pending_name = ""
#     for t in texts:
#         m = SECREF_RX.fullmatch(t.strip())
#         if m:
#             refs[m.group(1)] = pending_name
#         else:
#             pending_name = t
#             for m2 in SECREF_RX.finditer(t):
#                 refs[m2.group(1)] = SECREF_RX.sub("", t).strip()
#     return refs


# # ---------- output 3: tables of referenced sections + pin mapping ----------

# def collect_section_tables(doc, wanted):
#     counters = [0] * 9
#     current = None
#     found = {}

#     for child in doc.element.body.iterchildren():
#         if child.tag.endswith("}p"):
#             p = Paragraph(child, doc)
#             m = HEADING_RX.match(p.style.name)
#             if m and full_text(child).strip():
#                 lvl = int(m.group(1))
#                 counters[lvl - 1] += 1
#                 for i in range(lvl, 9):
#                     counters[i] = 0
#                 number = ".".join(str(c) for c in counters[:lvl])
#                 current = number if number in wanted else None
#                 if current:
#                     found[current] = {
#                         "title": " ".join(full_text(child).split()),
#                         "tables": []}
#         elif child.tag.endswith("}tbl") and current:
#             t = Table(child, doc)
#             rows = []
#             mnemonic = ""
#             for r in t.rows:
#                 cells = [" ".join(live_text(c._tc).split()) for c in r.cells]
#                 dedup = [cells[0]] + [c for i, c in enumerate(cells[1:], 1)
#                                       if c != cells[i - 1]]
#                 if any(dedup):
#                     if dedup[0].lower() == "mnemonic" and len(dedup) > 1:
#                         mnemonic = dedup[1]
#                     # pin mapping: insert pin name after the row label
#                     rows.append([dedup[0], pin_for(dedup[0])] + dedup[1:])
#             if rows:
#                 if not mnemonic and found[current]["tables"]:
#                     mnemonic = found[current]["tables"][-1]["mnemonic"]
#                 found[current]["tables"].append(
#                     {"mnemonic": mnemonic, "rows": rows})
#     return found


# def write_excel(refs, sections, path):
#     from openpyxl import Workbook
#     from openpyxl.styles import Font, PatternFill

#     wb = Workbook()
#     bold = Font(bold=True)
#     pinfill = PatternFill("solid", start_color="DDEBF7")

#     ws = wb.active
#     ws.title = "Summary"
#     ws.append(["Section", "Section title", "Diagram box", "Tables found"])
#     for c in ws[1]:
#         c.font = bold
#     for sec in sorted(refs, key=lambda s: [int(x) for x in s.split(".")]):
#         info = sections.get(sec, {})
#         ws.append([sec, info.get("title", "NOT FOUND"),
#                    refs[sec], len(info.get("tables", []))])
#     for col, w in zip("ABCD", (10, 42, 45, 12)):
#         ws.column_dimensions[col].width = w

#     # ---- dedicated "Pins" sheet with every mapped pin ----
#     ps = wb.create_sheet("Pins")
#     ps.append(["Section", "Mnemonic", "Row label", "Pin",
#                "Value / condition from narrative"])
#     for c in ps[1]:
#         c.font = bold
#     for sec in sorted(sections, key=lambda s: [int(x) for x in s.split(".")]):
#         for tbl in sections[sec]["tables"]:
#             mnemonic = tbl["mnemonic"]
#             for row in tbl["rows"]:
#                 if len(row) > 1 and row[1]:      # mapped pin rows only
#                     value = row[2] if len(row) > 2 else ""
#                     ps.append([sec, mnemonic, row[0], row[1], value])
#                     ps.cell(ps.max_row, 4).fill = pinfill
#     for col, w in zip("ABCDE", (10, 30, 24, 10, 80)):
#         ps.column_dimensions[col].width = w

#     # ---- one sheet per section (full tables) ----
#     for sec in sorted(sections, key=lambda s: [int(x) for x in s.split(".")]):
#         info = sections[sec]
#         sh = wb.create_sheet(("Sec " + sec)[:31])
#         sh.append(["Section %s - %s" % (sec, info["title"])])
#         sh["A1"].font = bold
#         for n, tbl in enumerate(info["tables"], 1):
#             sh.append([])
#             sh.append(["Table %d (%s)" % (n, tbl["mnemonic"] or "-"),
#                        "Pin", "Value"])
#             for c in sh[sh.max_row]:
#                 c.font = bold
#             for row in tbl["rows"]:
#                 sh.append(row)
#                 if len(row) > 1 and row[1]:          # highlight mapped pins
#                     sh.cell(sh.max_row, 2).fill = pinfill
#         sh.column_dimensions["A"].width = 32
#         sh.column_dimensions["B"].width = 12
#         for col in "CDEF":
#             sh.column_dimensions[col].width = 50

#     wb.save(path)


# def write_mnemonic_workbooks(sections, out_root):
#     """One folder per mnemonic, each with <Mnemonic>.xlsx containing
#     that mnemonic's table sheet and its own Pins sheet."""
#     from openpyxl import Workbook
#     from openpyxl.styles import Font, PatternFill
#     bold = Font(bold=True)
#     pinfill = PatternFill("solid", start_color="DDEBF7")

#     # group tables by mnemonic
#     per = {}
#     for sec in sorted(sections, key=lambda s: [int(x) for x in s.split(".")]):
#         info = sections[sec]
#         for tbl in info["tables"]:
#             mn = tbl["mnemonic"] or "NoMnemonic"
#             per.setdefault(mn, []).append((sec, info["title"], tbl))

#     for mn, items in per.items():
#         folder = out_root / mn
#         folder.mkdir(parents=True, exist_ok=True)
#         wb = Workbook()

#         # sheet 1: the mnemonic's tables
#         sh = wb.active
#         sh.title = mn[:31]
#         for sec, title, tbl in items:
#             sh.append(["Section %s - %s" % (sec, title)])
#             sh.cell(sh.max_row, 1).font = bold
#             sh.append(["Row label", "Pin", "Value"])
#             for c in sh[sh.max_row]:
#                 c.font = bold
#             for row in tbl["rows"]:
#                 sh.append(row)
#                 if len(row) > 1 and row[1]:
#                     sh.cell(sh.max_row, 2).fill = pinfill
#             sh.append([])
#         sh.column_dimensions["A"].width = 32
#         sh.column_dimensions["B"].width = 12
#         for col in "CDEF":
#             sh.column_dimensions[col].width = 60

#         # sheet 2: pins of this mnemonic only
#         ps = wb.create_sheet("Pins")
#         ps.append(["Section", "Row label", "Pin",
#                    "Value / condition from narrative"])
#         for c in ps[1]:
#             c.font = bold
#         for sec, title, tbl in items:
#             for row in tbl["rows"]:
#                 if len(row) > 1 and row[1]:
#                     value = row[2] if len(row) > 2 else ""
#                     ps.append([sec, row[0], row[1], value])
#                     ps.cell(ps.max_row, 3).fill = pinfill
#         for col, w in zip("ABCD", (10, 24, 10, 90)):
#             ps.column_dimensions[col].width = w

#         wb.save(folder / ("%s.xlsx" % mn))
#         print("  %s -> %s" % (mn, folder / ("%s.xlsx" % mn)))


# # ---------- main ----------

# def main():
#     print("Document : %s" % DOC_PATH)
#     if not Path(DOC_PATH).exists():
#         raise SystemExit("Document not found! Check DOC_PATH.")

#     out_root = Path(OUTPUT_DIR)
#     out_dir = out_root / "overall"          # combined results live here
#     out_dir.mkdir(parents=True, exist_ok=True)

#     doc = docx.Document(DOC_PATH)
#     image_rid, ole_rid = find_section_objects(doc)

#     print("\n--- 1/3 IMAGE ---")
#     if image_rid:
#         save_image(doc, image_rid, out_dir)
#     else:
#         print("No preview image found.")

#     print("\n--- 2/3 DIAGRAM TEXT (not struck out) ---")
#     refs = {}
#     if ole_rid:
#         part = doc.part.related_parts[ole_rid]
#         texts = diagram_kept_texts(part.blob)
#         for t in texts:
#             print("  " + t)
#         (out_dir / TEXT_NAME).write_text("\n".join(texts), encoding="utf-8")
#         print("Saved: %s" % (out_dir / TEXT_NAME))
#         refs = pair_refs_with_names(texts)
#         refs.pop(TARGET_SECTION, None)
#     else:
#         print("No embedded Visio object found - skipping text/tables.")

#     print("\n--- 3/3 REFERENCED SECTION TABLES (with pin mapping) ---")
#     if refs:
#         print("References: %s" % ", ".join(sorted(refs)))
#         sections = collect_section_tables(doc, set(refs))
#         mapped = unmapped = 0
#         for sec, info in sections.items():
#             for tbl in info["tables"]:
#                 for row in tbl["rows"]:
#                     if row[1]:
#                         mapped += 1
#                     elif row[0]:
#                         unmapped += 1
#         for sec in sorted(refs, key=lambda s: [int(x) for x in s.split(".")]):
#             info = sections.get(sec)
#             if info:
#                 print("  Sec %-7s %-45s -> %d table(s)"
#                       % (sec, info["title"][:45], len(info["tables"])))
#             else:
#                 print("  Sec %-7s NOT FOUND in document" % sec)
#         print("Pin mapping: %d row(s) mapped, %d without a pin" % (mapped, unmapped))
#         write_excel(refs, sections, out_dir / XLSX_NAME)
#         print("Saved: %s" % (out_dir / XLSX_NAME))

#         print("\n--- PER-MNEMONIC FOLDERS ---")
#         write_mnemonic_workbooks(sections, out_root)

#     print("\nDONE - overall results: %s ; per-mnemonic folders next to it." % out_dir)


# if __name__ == "__main__":
#     try:
#         main()
#     except SystemExit as e:
#         print("STOPPED:", e)
#     except Exception:
#         traceback.print_exc()
#     input("Press Enter to close...")

############################folder + Image exctraction############
"""
ALL-IN-ONE extractor for the HRSG control narrative (.docx):

  Output 1: the block diagram of section 4.1.1 as PNG (via PowerPoint)
            (OFF by default - see EXTRACT_DIAGRAM_IMAGE in CONFIG)
  Output 2: the not-struck-out diagram text  -> .txt
  Output 3: tables of all sections referenced by the diagram -> .xlsx
            with PIN MAPPING: every table row label is mapped to its
            MarkVIe actuator block pin name (extra "Pin" column)
            plus a dedicated "Pins" sheet listing all mapped pins

  Folder structure of the output:
    OUTPUT_DIR\\overall\\      -> text + combined Excel
    OUTPUT_DIR\\<Mnemonic>\\   -> <Mnemonic>.xlsx (its sheet + Pins sheet)
                                 + images of sections referenced in the
                                   pin values, named
                                   <Mnemonic>__<PIN>__Sec_<x.y.z>.png

Requires:  pip install python-docx openpyxl pillow pywin32
Usage   :  edit CONFIG, then run  python Image_extract.py
"""

print("script started")

import io
import re
import traceback
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

# ============================================================
#                          CONFIG
# ============================================================
# DOC_PATH = r"C:\Users\1052088\Downloads\GEV\YTL-50-E-LBA50---EN-LO-001 - 1GP052419 - A-en-1GP052419AT03-en.docx"
DOC_PATH = r"D:\Project\Mallikarjun\Rajat_Task\backend-gev\YTL-50-E-LBA50---EN-LO-001 - 1GP052419 - A-en-1GP052419AT03-en.docx"

TARGET_SECTION = "4.1.1"          # section containing the block diagram

OUTPUT_DIR = r"D:\Project\Mallikarjun\Rajat_Task\backend-gev\output\out_image_1"
IMAGE_NAME = "section_4_1_1_diagram"          # -> .emf + .png
TEXT_NAME  = "section_4_1_1_diagram_text.txt"
XLSX_NAME  = "referenced_section_tables.xlsx"

PNG_SCALE = 4                     # enlarge before export = sharper PNG

EXTRACT_DIAGRAM_IMAGE    = False  # section 4.1.1 drawing in overall folder
EXTRACT_REFERENCE_IMAGES = True   # per-pin section images in mnemonic folders

# Narrative table row label  ->  MarkVIe block pin name
# (add more entries here any time - matching is case-insensitive)
PIN_MAP = {
    # group (ON/OFF) pins
    "set auto mode":      "AU_SEL",
    "permit on":          "ON_PMT",
    "auto command on":    "AU_ON",
    "force on":           "ON_FRC",
    "permit off":         "OFF_PMT",
    "auto command off":   "AU_OFF",
    "force off":          "OFF_FRC",
    # valve (OPEN/CLOSE) pins
    "permit open":        "OP_PMT",
    "auto command open":  "AU_OP",
    "force open":         "OP_FRC",
    "override open":      "OP_OVR",
    "permit close":       "CL_PMT",
    "auto command close": "AU_CL",
    "force close":        "CL_FRC",
    "override close":     "CL_OVR",
    "auto permit":        "AU_PMT",
    "force condition":    "CMD_FRC",
    "force command":      "CMD_FRC",
    # control valve pins
    "modulate permit":    "MOD_PMT",
    "modulate value":     "RSP",
    "force value":        "FV",
    "auto setpoint":      "AU_SETP",
}
# ============================================================

VNS = "{http://schemas.microsoft.com/office/visio/2012/main}"
WNS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEADING_RX = re.compile(r"^Heading (\d)$")
SECREF_RX = re.compile(r"\[\s*\u00a7\s*([\d.]+)\s*\]")


def full_text(el):
    """Text including tracked-change deletions (original document)."""
    return "".join(n.text or "" for n in el.iter()
                   if n.tag in (WNS + "t", WNS + "delText"))


def live_text(el):
    """Text as it reads after accepting tracked changes."""
    return "".join(n.text or "" for n in el.iter() if n.tag == WNS + "t")


def pin_for(label):
    """Map a narrative row label to a block pin name ('' if unknown)."""
    return PIN_MAP.get(" ".join(label.lower().split()), "")


# ---------- locate diagram objects in the target section ----------

def find_section_objects(doc):
    """Return (image_rid, ole_rid) of the diagram in TARGET_SECTION."""
    counters = [0] * 9
    depth = TARGET_SECTION.count(".") + 1
    inside = False
    image_rid = ole_rid = None

    for par in doc.paragraphs:
        m = HEADING_RX.match(par.style.name)
        if m and full_text(par._p).strip():
            lvl = int(m.group(1))
            counters[lvl - 1] += 1
            for i in range(lvl, 9):
                counters[i] = 0
            number = ".".join(str(c) for c in counters[:lvl])
            if inside and lvl <= depth:
                break
            if number == TARGET_SECTION:
                inside = True
                print("Found section %s: %s" % (number, par.text.strip()))
            continue
        if inside:
            xml = par._p.xml
            if image_rid is None:
                m2 = re.search(r'<v:imagedata[^>]*r:id="(rId\d+)"', xml)
                if m2:
                    image_rid = m2.group(1)
            if ole_rid is None:
                m3 = re.search(r'<o:OLEObject[^>]*r:id="(rId\d+)"', xml)
                if m3:
                    ole_rid = m3.group(1)
            if image_rid and ole_rid:
                break

    if not (image_rid or ole_rid):
        raise SystemExit("No diagram found in section %s." % TARGET_SECTION)
    return image_rid, ole_rid


# ---------- output 1: image ----------

def save_image(doc, image_rid, out_dir):
    part = doc.part.related_parts[image_rid]
    ext = Path(str(part.partname)).suffix
    raw = out_dir / (IMAGE_NAME + ext)
    raw.write_bytes(part.blob)
    print("Saved original image: %s" % raw)

    if ext.lower() not in (".emf", ".wmf"):
        return

    png = out_dir / (IMAGE_NAME + ".png")
    try:
        import win32com.client
        print("Starting PowerPoint for EMF -> PNG conversion ...")
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt.Presentations.Add(WithWindow=0)
        try:
            slide = pres.Slides.Add(1, 12)
            shape = slide.Shapes.AddPicture(
                FileName=str(raw), LinkToFile=0, SaveWithDocument=-1,
                Left=0, Top=0)
            shape.ScaleWidth(PNG_SCALE, 1)
            shape.ScaleHeight(PNG_SCALE, 1)
            shape.Export(str(png), 2)          # 2 = PNG
        finally:
            pres.Close()
            ppt.Quit()

        try:
            from PIL import Image, ImageChops
            img = Image.open(png)
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                white = Image.new("RGBA", img.size,
                                  (255, 255, 255, 255))
                img = Image.alpha_composite(white, img).convert("RGB")
            else:
                img = img.convert("RGB")
            bg = Image.new("RGB", img.size, (255, 255, 255))
            bbox = ImageChops.difference(img, bg).getbbox()
            if bbox:
                img.crop(bbox).save(png)
        except ImportError:
            pass
        print("Saved PNG           : %s" % png)
    except Exception as e:
        print("[!] PNG conversion failed (%s); original %s kept." % (e, ext))


# ---------- output 2: diagram text (not struck out) ----------

def diagram_kept_texts(vsdx_bytes):
    zf = zipfile.ZipFile(io.BytesIO(vsdx_bytes))
    page = [n for n in zf.namelist()
            if re.match(r"visio/pages/page\d+\.xml$", n)][0]
    root = ET.fromstring(zf.read(page))

    out = []
    for shape in root.iter(VNS + "Shape"):
        text_el = shape.find(VNS + "Text")
        if text_el is None:
            continue
        struck_rows = {}
        for sec in shape.findall(VNS + "Section"):
            if sec.get("N") == "Character":
                for row in sec.findall(VNS + "Row"):
                    ix = int(row.get("IX", "0"))
                    s = any(c.get("N") == "Strikethru" and c.get("V") == "1"
                            for c in row.findall(VNS + "Cell"))
                    struck_rows[ix] = s
        kept = []
        state = {"cur": 0}

        def add(t):
            if t and not struck_rows.get(state["cur"], False):
                kept.append(t)

        add(text_el.text)
        for ch in text_el:
            if ch.tag == VNS + "cp":
                state["cur"] = int(ch.get("IX", "0"))
            add(ch.tail)
        txt = " ".join("".join(kept).split())
        if txt:
            out.append(txt)
    return out


def pair_refs_with_names(texts):
    refs = {}
    pending_name = ""
    for t in texts:
        m = SECREF_RX.fullmatch(t.strip())
        if m:
            refs[m.group(1)] = pending_name
        else:
            pending_name = t
            for m2 in SECREF_RX.finditer(t):
                refs[m2.group(1)] = SECREF_RX.sub("", t).strip()
    return refs


# ---------- output 3: tables of referenced sections + pin mapping ----------

def collect_section_tables(doc, wanted):
    counters = [0] * 9
    current = None
    found = {}

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p = Paragraph(child, doc)
            m = HEADING_RX.match(p.style.name)
            if m and full_text(child).strip():
                lvl = int(m.group(1))
                counters[lvl - 1] += 1
                for i in range(lvl, 9):
                    counters[i] = 0
                number = ".".join(str(c) for c in counters[:lvl])
                current = number if number in wanted else None
                if current:
                    found[current] = {
                        "title": " ".join(full_text(child).split()),
                        "tables": []}
        elif child.tag.endswith("}tbl") and current:
            t = Table(child, doc)
            rows = []
            mnemonic = ""
            for r in t.rows:
                cells = [" ".join(live_text(c._tc).split()) for c in r.cells]
                dedup = [cells[0]] + [c for i, c in enumerate(cells[1:], 1)
                                      if c != cells[i - 1]]
                if any(dedup):
                    if dedup[0].lower() == "mnemonic" and len(dedup) > 1:
                        mnemonic = dedup[1]
                    # pin mapping: insert pin name after the row label
                    rows.append([dedup[0], pin_for(dedup[0])] + dedup[1:])
            if rows:
                if not mnemonic and found[current]["tables"]:
                    mnemonic = found[current]["tables"][-1]["mnemonic"]
                found[current]["tables"].append(
                    {"mnemonic": mnemonic, "rows": rows})
    return found


def write_excel(refs, sections, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    pinfill = PatternFill("solid", start_color="DDEBF7")

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Section", "Section title", "Diagram box", "Tables found"])
    for c in ws[1]:
        c.font = bold
    for sec in sorted(refs, key=lambda s: [int(x) for x in s.split(".")]):
        info = sections.get(sec, {})
        ws.append([sec, info.get("title", "NOT FOUND"),
                   refs[sec], len(info.get("tables", []))])
    for col, w in zip("ABCD", (10, 42, 45, 12)):
        ws.column_dimensions[col].width = w

    # ---- dedicated "Pins" sheet with every mapped pin ----
    ps = wb.create_sheet("Pins")
    ps.append(["Section", "Mnemonic", "Row label", "Pin",
               "Value / condition from narrative"])
    for c in ps[1]:
        c.font = bold
    for sec in sorted(sections, key=lambda s: [int(x) for x in s.split(".")]):
        for tbl in sections[sec]["tables"]:
            mnemonic = tbl["mnemonic"]
            for row in tbl["rows"]:
                if len(row) > 1 and row[1]:      # mapped pin rows only
                    value = row[2] if len(row) > 2 else ""
                    ps.append([sec, mnemonic, row[0], row[1], value])
                    ps.cell(ps.max_row, 4).fill = pinfill
    for col, w in zip("ABCDE", (10, 30, 24, 10, 80)):
        ps.column_dimensions[col].width = w

    # ---- one sheet per section (full tables) ----
    for sec in sorted(sections, key=lambda s: [int(x) for x in s.split(".")]):
        info = sections[sec]
        sh = wb.create_sheet(("Sec " + sec)[:31])
        sh.append(["Section %s - %s" % (sec, info["title"])])
        sh["A1"].font = bold
        for n, tbl in enumerate(info["tables"], 1):
            sh.append([])
            sh.append(["Table %d (%s)" % (n, tbl["mnemonic"] or "-"),
                       "Pin", "Value"])
            for c in sh[sh.max_row]:
                c.font = bold
            for row in tbl["rows"]:
                sh.append(row)
                if len(row) > 1 and row[1]:          # highlight mapped pins
                    sh.cell(sh.max_row, 2).fill = pinfill
        sh.column_dimensions["A"].width = 32
        sh.column_dimensions["B"].width = 12
        for col in "CDEF":
            sh.column_dimensions[col].width = 50

    wb.save(path)


def write_mnemonic_workbooks(sections, out_root):
    """One folder per mnemonic, each with <Mnemonic>.xlsx containing
    that mnemonic's table sheet and its own Pins sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    bold = Font(bold=True)
    pinfill = PatternFill("solid", start_color="DDEBF7")

    # group tables by mnemonic
    per = {}
    for sec in sorted(sections, key=lambda s: [int(x) for x in s.split(".")]):
        info = sections[sec]
        for tbl in info["tables"]:
            mn = tbl["mnemonic"] or "NoMnemonic"
            per.setdefault(mn, []).append((sec, info["title"], tbl))

    for mn, items in per.items():
        folder = out_root / mn
        folder.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        # sheet 1: the mnemonic's tables
        sh = wb.active
        sh.title = mn[:31]
        for sec, title, tbl in items:
            sh.append(["Section %s - %s" % (sec, title)])
            sh.cell(sh.max_row, 1).font = bold
            sh.append(["Row label", "Pin", "Value"])
            for c in sh[sh.max_row]:
                c.font = bold
            for row in tbl["rows"]:
                sh.append(row)
                if len(row) > 1 and row[1]:
                    sh.cell(sh.max_row, 2).fill = pinfill
            sh.append([])
        sh.column_dimensions["A"].width = 32
        sh.column_dimensions["B"].width = 12
        for col in "CDEF":
            sh.column_dimensions[col].width = 60

        # sheet 2: pins of this mnemonic only
        ps = wb.create_sheet("Pins")
        ps.append(["Section", "Row label", "Pin",
                   "Value / condition from narrative"])
        for c in ps[1]:
            c.font = bold
        for sec, title, tbl in items:
            for row in tbl["rows"]:
                if len(row) > 1 and row[1]:
                    value = row[2] if len(row) > 2 else ""
                    ps.append([sec, row[0], row[1], value])
                    ps.cell(ps.max_row, 3).fill = pinfill
        for col, w in zip("ABCD", (10, 24, 10, 90)):
            ps.column_dimensions[col].width = w

        wb.save(folder / ("%s.xlsx" % mn))
        print("  %s -> %s" % (mn, folder / ("%s.xlsx" % mn)))


# ---------- reference images: [Sec x.y.z] inside pin values ----------

def collect_pin_section_refs(sections):
    """Scan every mnemonic table; return [(mnemonic, pin, ref_section)]."""
    refs = []
    seen = set()
    for sec in sections:
        for tbl in sections[sec]["tables"]:
            mn = tbl["mnemonic"] or "NoMnemonic"
            for row in tbl["rows"]:
                if len(row) > 1 and row[1]:            # mapped pin rows
                    value = " ".join(str(c) for c in row[2:])
                    for m in SECREF_RX.finditer(value):
                        key = (mn, row[1], m.group(1))
                        if key not in seen:
                            seen.add(key)
                            refs.append(key)
    return refs


def map_section_images(doc):
    """Walk the document once: section number -> [image rIds]."""
    counters = [0] * 9
    current = None
    imgs = {}
    for par in doc.paragraphs:
        m = HEADING_RX.match(par.style.name)
        if m and full_text(par._p).strip():
            lvl = int(m.group(1))
            counters[lvl - 1] += 1
            for i in range(lvl, 9):
                counters[i] = 0
            current = ".".join(str(c) for c in counters[:lvl])
            continue
        if current:
            xml = par._p.xml
            m2 = (re.search(r'<v:imagedata[^>]*r:id="(rId\d+)"', xml)
                  or re.search(r'r:embed="(rId\d+)"', xml))
            if m2:
                imgs.setdefault(current, []).append(m2.group(1))
    return imgs


def extract_reference_images(doc, sections, out_root):
    """For every [Sec x.y.z] found in a pin's value, extract that
    section's image(s) into the mnemonic folder, named
    <Mnemonic>__<PIN>__Sec_<x.y.z>.png (original .emf kept too)."""
    refs = collect_pin_section_refs(sections)
    if not refs:
        print("  no section references found in pin values")
        return
    sec_imgs = map_section_images(doc)

    emf_files = []
    for mn, pin, refsec in refs:
        rids = sec_imgs.get(refsec)
        if not rids:
            print("  %-28s %-8s Sec %-7s -> no image in that section"
                  % (mn, pin, refsec))
            continue
        folder = out_root / mn
        folder.mkdir(parents=True, exist_ok=True)
        for n, rid in enumerate(rids, 1):
            part = doc.part.related_parts[rid]
            ext = Path(str(part.partname)).suffix
            suffix = "" if len(rids) == 1 else "_%d" % n
            fname = "%s__%s__Sec_%s%s%s" % (mn, pin, refsec, suffix, ext)
            fpath = folder / fname
            fpath.write_bytes(part.blob)
            print("  %-28s %-8s Sec %-7s -> %s" % (mn, pin, refsec, fname))
            if ext.lower() in (".emf", ".wmf"):
                emf_files.append(fpath)

    # convert all extracted EMFs to PNG (per-file error handling)
    if emf_files:
        ok = fail = 0
        try:
            import win32com.client
            print("  Converting %d image(s) to PNG via PowerPoint ..."
                  % len(emf_files))
            ppt = win32com.client.Dispatch("PowerPoint.Application")
            pres = ppt.Presentations.Add(WithWindow=0)
            try:
                for emf in emf_files:
                    try:
                        slide = pres.Slides.Add(1, 12)
                        shape = slide.Shapes.AddPicture(
                            FileName=str(emf), LinkToFile=0,
                            SaveWithDocument=-1, Left=0, Top=0)
                        shape.ScaleWidth(PNG_SCALE, 1)
                        shape.ScaleHeight(PNG_SCALE, 1)
                        png = emf.with_suffix(".png")
                        shape.Export(str(png), 2)
                        slide.Delete()
                        try:
                            from PIL import Image, ImageChops
                            img = Image.open(png)
                            if img.mode in ("RGBA", "LA", "P"):
                                img = img.convert("RGBA")
                                white = Image.new(
                                    "RGBA", img.size,
                                    (255, 255, 255, 255))
                                img = Image.alpha_composite(
                                    white, img).convert("RGB")
                            else:
                                img = img.convert("RGB")
                            bg = Image.new("RGB", img.size,
                                           (255, 255, 255))
                            bbox = ImageChops.difference(
                                img, bg).getbbox()
                            if bbox:
                                img.crop(bbox).save(png)
                        except ImportError:
                            pass
                        ok += 1
                        print("    OK   %s" % png.name)
                    except Exception as e:
                        fail += 1
                        print("    FAIL %s (%s)" % (emf.name, e))
            finally:
                pres.Close()
                ppt.Quit()
            print("  PNG conversion: %d ok, %d failed." % (ok, fail))
        except Exception as e:
            print("  [!] PowerPoint not available (%s); "
                  ".emf files kept." % e)


# ---------- main ----------

def main():
    print("Document : %s" % DOC_PATH)
    if not Path(DOC_PATH).exists():
        raise SystemExit("Document not found! Check DOC_PATH.")

    out_root = Path(OUTPUT_DIR)
    out_dir = out_root / "overall"          # combined results live here
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = docx.Document(DOC_PATH)
    image_rid, ole_rid = find_section_objects(doc)

    print("\n--- 1/3 IMAGE ---")
    if not EXTRACT_DIAGRAM_IMAGE:
        print("skipped (EXTRACT_DIAGRAM_IMAGE = False)")
    elif image_rid:
        save_image(doc, image_rid, out_dir)
    else:
        print("No preview image found.")

    print("\n--- 2/3 DIAGRAM TEXT (not struck out) ---")
    refs = {}
    if ole_rid:
        part = doc.part.related_parts[ole_rid]
        texts = diagram_kept_texts(part.blob)
        for t in texts:
            print("  " + t)
        (out_dir / TEXT_NAME).write_text("\n".join(texts), encoding="utf-8")
        print("Saved: %s" % (out_dir / TEXT_NAME))
        refs = pair_refs_with_names(texts)
        refs.pop(TARGET_SECTION, None)
    else:
        print("No embedded Visio object found - skipping text/tables.")

    print("\n--- 3/3 REFERENCED SECTION TABLES (with pin mapping) ---")
    if refs:
        print("References: %s" % ", ".join(sorted(refs)))
        sections = collect_section_tables(doc, set(refs))
        mapped = unmapped = 0
        for sec, info in sections.items():
            for tbl in info["tables"]:
                for row in tbl["rows"]:
                    if row[1]:
                        mapped += 1
                    elif row[0]:
                        unmapped += 1
        for sec in sorted(refs, key=lambda s: [int(x) for x in s.split(".")]):
            info = sections.get(sec)
            if info:
                print("  Sec %-7s %-45s -> %d table(s)"
                      % (sec, info["title"][:45], len(info["tables"])))
            else:
                print("  Sec %-7s NOT FOUND in document" % sec)
        print("Pin mapping: %d row(s) mapped, %d without a pin" % (mapped, unmapped))
        write_excel(refs, sections, out_dir / XLSX_NAME)
        print("Saved: %s" % (out_dir / XLSX_NAME))

        print("\n--- PER-MNEMONIC FOLDERS ---")
        write_mnemonic_workbooks(sections, out_root)

        if EXTRACT_REFERENCE_IMAGES:
            print("\n--- REFERENCED-SECTION IMAGES (per pin) ---")
            extract_reference_images(doc, sections, out_root)
        else:
            print("\n--- REFERENCED-SECTION IMAGES skipped "
                  "(EXTRACT_REFERENCE_IMAGES = False) ---")

    print("\nDONE - overall results: %s ; per-mnemonic folders next to it." % out_dir)


if __name__ == "__main__":
    try:
        main()
    except SystemExit as e:
        print("STOPPED:", e)
    except Exception:
        traceback.print_exc()
    input("Press Enter to close...")
